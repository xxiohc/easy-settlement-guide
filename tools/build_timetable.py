#!/usr/bin/env python3
"""KORAIL KTX 시간표 xlsx → data/ktx_timetable.json

시트마다 '열차번호' 헤더가 하행/상행 두 블록으로 들어 있고,
역 이름이 가로로, 열차가 세로로 놓인 표다. 00:00:00 은 해당 역 미정차를 뜻한다.
"""
import json, re, sys
from datetime import time as dtime, datetime
from pathlib import Path

import openpyxl

APP = Path(__file__).resolve().parent.parent
SRC = APP / "data" / "source" / "KTX 시간표(202610 기준).xlsx"
OUT = APP / "data" / "ktx_timetable.json"

DAY_MAP = {
    "매일": [0, 1, 2, 3, 4, 5, 6],
    "월": [0], "화": [1], "수": [2], "목": [3], "금": [4], "토": [5], "일": [6],
}


def to_minutes(v):
    if v is None:
        return None
    if isinstance(v, (dtime, datetime)):
        return v.hour * 60 + v.minute
    s = str(v).strip()
    m = re.match(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?$", s)
    if not m:
        return None
    return int(m.group(1)) * 60 + int(m.group(2))


def parse_days(note):
    if not note:
        return DAY_MAP["매일"]
    s = str(note).strip()
    if "매일" in s:
        return DAY_MAP["매일"]
    days = sorted({d for ch in s if ch in DAY_MAP for d in DAY_MAP[ch]})
    return days or DAY_MAP["매일"]


def grid(ws):
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


def find_blocks(rows):
    """('열차번호' 셀 위치) → (헤더행, 열차번호열, [(역명, 열)], 비고열)"""
    blocks = []
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            if str(val).strip() != "열차번호":
                continue
            stations, note_col = [], None
            cc = c + 1
            while cc < len(row):
                name = str(row[cc]).strip() if row[cc] is not None else ""
                if name == "비고":
                    note_col = cc
                    break
                if name == "":
                    break
                if name != "편성":
                    stations.append((name, cc))
                cc += 1
            if len(stations) >= 3:
                blocks.append((r, c, stations, note_col))
    return blocks


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    trains, seen = [], {}
    for sheet in wb.sheetnames:
        rows = grid(wb[sheet])
        for hdr, tno_col, stations, note_col in find_blocks(rows):
            for r in range(hdr + 1, len(rows)):
                row = rows[r]
                if tno_col >= len(row):
                    continue
                raw = row[tno_col]
                if raw is None:
                    continue
                tno = str(raw).strip()
                if not re.fullmatch(r"\d{1,5}", tno):
                    continue
                ttype = str(row[tno_col + 1] or "").strip()
                stops, prev, carry = [], None, 0
                for name, col in stations:
                    if col >= len(row):
                        continue
                    mins = to_minutes(row[col])
                    if mins is None or mins == 0:
                        continue
                    if prev is not None and mins + carry < prev:
                        carry += 1440
                    mins += carry
                    stops.append({"s": name, "t": mins})
                    prev = mins
                if len(stops) < 2:
                    continue
                note = row[note_col] if note_col is not None and note_col < len(row) else None
                key = (tno, stops[0]["s"], stops[0]["t"])
                if key in seen:
                    continue
                seen[key] = True
                trains.append({
                    "no": tno,
                    "type": ttype or "KTX",
                    "line": sheet,
                    "days": parse_days(note),
                    "note": str(note).strip() if note else "매일",
                    "stops": stops,
                })

    stations = sorted({st["s"] for t in trains for st in t["stops"]})
    out = {
        "source": SRC.name,
        "basis": "KORAIL KTX 시간표 (2026년 10월 기준)",
        "note": "t = 00:00부터의 분(초 단위는 버림 — 출발편을 늦게 보지 않기 위함). 1440 이상은 익일. 00:00:00 셀은 미정차로 처리.",
        "updatedAt": datetime.now().strftime("%Y-%m-%d"),
        "trainCount": len(trains),
        "stationCount": len(stations),
        "stations": stations,
        "trains": trains,
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"열차 {len(trains)}편 · 역 {len(stations)}개 → {OUT}")
    masan = [t for t in trains if any(s["s"] == "마산" for s in t["stops"])]
    print(f"마산 정차 {len(masan)}편")


if __name__ == "__main__":
    sys.exit(main())
